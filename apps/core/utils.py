"""
Utilitários para o sistema HR, incluindo geração de documentos Word e PDF.
"""
import os
from datetime import datetime
from decimal import Decimal
from django.conf import settings
from django.core.mail import EmailMessage
from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors


def format_currency(value):
    """
    Formata valor monetário para exibição.
    """
    if value is None:
        return "0,00 MT"
    return f"{value:,.2f} MT".replace(',', 'X').replace('.', ',').replace('X', '.')


def generate_word_declaration(employee, template_path=None):
    """
    Gera declaração de rendimentos em formato Word usando python-docx.
    
    Args:
        employee: Instância do modelo Employee
        template_path: Caminho para template Word (opcional)
    
    Returns:
        Caminho do arquivo gerado
    """
    # Criar documento Word
    if template_path and os.path.exists(template_path):
        doc = Document(template_path)
    else:
        doc = Document()
    
    # Configurar margens
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Cabeçalho da empresa
    header = doc.add_heading('INFRASECUR MOÇAMBIQUE LTD', 0)
    header.alignment = 1  # Centralizado
    
    company_info = doc.add_paragraph()
    company_info.alignment = 1
    company_info.add_run('Rua da Resistência, nº 1234\n')
    company_info.add_run('Maputo, Moçambique\n')
    company_info.add_run('Tel: +258 21 123 456\n')
    company_info.add_run('Email: rh@infrasecur.co.mz\n')
    
    # Espaço
    doc.add_paragraph()
    
    # Título do documento
    title = doc.add_heading('DECLARAÇÃO DE RENDIMENTOS', 1)
    title.alignment = 1
    
    # Data
    date_p = doc.add_paragraph()
    date_p.alignment = 2  # Alinhado à direita
    date_p.add_run(f'Maputo, {datetime.now().strftime("%d de %B de %Y")}')
    
    # Espaço
    doc.add_paragraph()
    
    # Corpo da declaração
    body = doc.add_paragraph()
    body.add_run('Declaramos para os devidos efeitos que o(a) Sr.(a) ')
    body.add_run(employee.full_name).bold = True
    body.add_run(f', portador(a) do documento de identificação nº ')
    body.add_run(employee.id_document).bold = True
    
    if employee.nuit:
        body.add_run(f', NUIT nº ')
        body.add_run(employee.nuit).bold = True
    
    body.add_run(f', exerce funções de ')
    body.add_run(employee.position).bold = True
    body.add_run(' nesta empresa desde ')
    body.add_run(employee.start_date.strftime('%d/%m/%Y')).bold = True
    body.add_run('.')
    
    # Informações salariais
    salary_p = doc.add_paragraph()
    salary_p.add_run('O referido colaborador aufere mensalmente um salário líquido de ')
    salary_p.add_run(format_currency(employee.net_salary)).bold = True
    
    if employee.daily_allowance and employee.daily_allowance > 0:
        salary_p.add_run(', subsídio de alimentação de ')
        salary_p.add_run(format_currency(employee.daily_allowance)).bold = True
    
    if employee.additional_remuneration and employee.additional_remuneration > 0:
        salary_p.add_run(', e outras remunerações no valor de ')
        salary_p.add_run(format_currency(employee.additional_remuneration)).bold = True
    
    salary_p.add_run('.')
    
    # Tipo de contrato
    contract_p = doc.add_paragraph()
    contract_p.add_run(f'O tipo de contrato é: ')
    contract_p.add_run(employee.get_contract_type_display()).bold = True
    contract_p.add_run('.')
    
    # Status
    status_p = doc.add_paragraph()
    status_p.add_run('Status atual: ')
    status_p.add_run(employee.get_status_display()).bold = True
    status_p.add_run('.')
    
    # Espaço
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Finalização
    final_p = doc.add_paragraph()
    final_p.add_run('Esta declaração é emitida a pedido do interessado para os fins que julgar convenientes.')
    
    # Espaço para assinatura
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    signature_p = doc.add_paragraph()
    signature_p.alignment = 1
    signature_p.add_run('_' * 40)
    signature_p.add_run('\nDepartamento de Recursos Humanos\n')
    signature_p.add_run('INFRASECUR MOÇAMBIQUE LTD')
    
    # Salvar documento
    filename = f"declaracao_rendimentos_{employee.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    file_path = os.path.join(settings.MEDIA_ROOT, 'documents', 'declarations', filename)
    
    # Criar diretório se não existir
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    
    doc.save(file_path)
    
    return file_path


def generate_pdf_declaration(employee):
    """
    Gera declaração de rendimentos em formato PDF usando ReportLab (fallback).
    
    Args:
        employee: Instância do modelo Employee
    
    Returns:
        Caminho do arquivo gerado
    """
    filename = f"declaracao_rendimentos_{employee.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    file_path = os.path.join(settings.MEDIA_ROOT, 'documents', 'declarations', filename)
    
    # Criar diretório se não existir
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    
    # Criar documento PDF
    doc = SimpleDocTemplate(file_path, pagesize=A4)
    styles = getSampleStyleSheet()
    
    # Estilo personalizado para título
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        spaceAfter=30,
        alignment=1  # Centralizado
    )
    
    # Estilo para cabeçalho da empresa
    company_style = ParagraphStyle(
        'CompanyHeader',
        parent=styles['Normal'],
        fontSize=14,
        alignment=1,
        spaceAfter=20
    )
    
    # Conteúdo do documento
    story = []
    
    # Cabeçalho da empresa
    story.append(Paragraph('<b>INFRASECUR MOÇAMBIQUE LTD</b>', company_style))
    story.append(Paragraph('Rua da Resistência, nº 1234<br/>Maputo, Moçambique<br/>Tel: +258 21 123 456<br/>Email: rh@infrasecur.co.mz', styles['Normal']))
    story.append(Spacer(1, 20))
    
    # Título
    story.append(Paragraph('<b>DECLARAÇÃO DE RENDIMENTOS</b>', title_style))
    
    # Data
    date_style = ParagraphStyle('DateStyle', parent=styles['Normal'], alignment=2)
    story.append(Paragraph(f'Maputo, {datetime.now().strftime("%d de %B de %Y")}', date_style))
    story.append(Spacer(1, 20))
    
    # Corpo da declaração
    body_text = f"""
    Declaramos para os devidos efeitos que o(a) Sr.(a) <b>{employee.full_name}</b>, 
    portador(a) do documento de identificação nº <b>{employee.id_document}</b>
    """
    
    if employee.nuit:
        body_text += f", NUIT nº <b>{employee.nuit}</b>"
    
    body_text += f"""
    , exerce funções de <b>{employee.position}</b> nesta empresa desde <b>{employee.start_date.strftime('%d/%m/%Y')}</b>.
    """
    
    story.append(Paragraph(body_text, styles['Normal']))
    story.append(Spacer(1, 12))
    
    # Informações salariais
    salary_text = f"O referido colaborador aufere mensalmente um salário líquido de <b>{format_currency(employee.net_salary)}</b>"
    
    if employee.daily_allowance and employee.daily_allowance > 0:
        salary_text += f", subsídio de alimentação de <b>{format_currency(employee.daily_allowance)}</b>"
    
    if employee.additional_remuneration and employee.additional_remuneration > 0:
        salary_text += f", e outras remunerações no valor de <b>{format_currency(employee.additional_remuneration)}</b>"
    
    salary_text += "."
    
    story.append(Paragraph(salary_text, styles['Normal']))
    story.append(Spacer(1, 12))
    
    # Tipo de contrato
    story.append(Paragraph(f"O tipo de contrato é: <b>{employee.get_contract_type_display()}</b>.", styles['Normal']))
    story.append(Spacer(1, 12))
    
    # Status
    story.append(Paragraph(f"Status atual: <b>{employee.get_status_display()}</b>.", styles['Normal']))
    story.append(Spacer(1, 20))
    
    # Finalização
    story.append(Paragraph("Esta declaração é emitida a pedido do interessado para os fins que julgar convenientes.", styles['Normal']))
    story.append(Spacer(1, 40))
    
    # Assinatura
    signature_style = ParagraphStyle('SignatureStyle', parent=styles['Normal'], alignment=1)
    story.append(Paragraph('_' * 40, signature_style))
    story.append(Paragraph('<b>Departamento de Recursos Humanos</b><br/>INFRASECUR MOÇAMBIQUE LTD', signature_style))
    
    # Gerar PDF
    doc.build(story)
    
    return file_path


def send_document_email(recipient_email, document_path, employee_name, document_type='Declaração de Rendimentos'):
    """
    Envia email com documento anexado.
    
    Args:
        recipient_email: Email do destinatário
        document_path: Caminho do documento a ser anexado
        employee_name: Nome do colaborador
        document_type: Tipo do documento
    
    Returns:
        Boolean indicando sucesso do envio
    """
    try:
        subject = f'{document_type} - {employee_name}'
        message = f"""
        Prezado(a) {employee_name},

        Segue em anexo o documento solicitado: {document_type}.

        Este documento foi gerado automaticamente pelo sistema HR da INFRASECUR MOÇAMBIQUE LTD.

        Atenciosamente,
        Departamento de Recursos Humanos
        INFRASECUR MOÇAMBIQUE LTD
        """
        
        email = EmailMessage(
            subject=subject,
            body=message,
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[recipient_email],
        )
        
        # Anexar documento
        if os.path.exists(document_path):
            email.attach_file(document_path)
        
        email.send()
        return True
        
    except Exception as e:
        print(f"Erro ao enviar email: {e}")
        return False


def get_user_permissions(user):
    """
    Retorna as permissões do usuário baseado nos grupos.
    
    Args:
        user: Instância do User
    
    Returns:
        Dict com permissões do usuário
    """
    permissions = {
        'is_admin': False,
        'is_employee': False,
        'can_manage_employees': False,
        'can_approve_documents': False,
        'can_moderate_knowledge': False,
        'can_manage_calendar': False,
    }
    
    if user.is_superuser:
        # Superuser tem todas as permissões
        for key in permissions:
            permissions[key] = True
        return permissions
    
    user_groups = user.groups.values_list('name', flat=True)
    
    if 'Admin' in user_groups:
        permissions.update({
            'is_admin': True,
            'can_manage_employees': True,
            'can_approve_documents': True,
            'can_moderate_knowledge': True,
            'can_manage_calendar': True,
        })
    
    if 'Colaborador' in user_groups:
        permissions['is_employee'] = True
    
    return permissions